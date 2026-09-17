import io
import json
import re

from docx import Document
from pypdf import PdfReader

from backend.core.settings import (
    Settings
)

from backend.services.openai_service import (
    OpenAIService
)


# ==========================================================
# EXCEPCIÓN
# ==========================================================

class DocumentServiceException(
    Exception
):

    def __init__(
        self,
        user_message: str,
        status_code: int = 400
    ):

        super().__init__(
            user_message
        )

        self.user_message = (
            user_message
        )

        self.status_code = (
            status_code
        )


# ==========================================================
# DOCUMENT SERVICE
# ==========================================================

class DocumentService:

    ALLOWED_EXTENSIONS = {
        "pdf",
        "docx",
        "txt"
    }


    ALLOWED_MIME_TYPES = {

        "pdf": {
            "application/pdf"
        },

        "docx": {
            (
                "application/"
                "vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            "application/octet-stream"
        },

        "txt": {
            "text/plain",
            "application/octet-stream"
        }

    }


    # Máximo de texto extraído que enviaremos a OpenAI.
    MAX_EXTRACTED_CHARACTERS = 30000


    # Fragmento máximo por sección.
    MAX_SECTION_CHARACTERS = 6000


    def __init__(
        self,
        settings: Settings,
        openai_service: OpenAIService
    ):

        self.settings = (
            settings
        )

        self.openai_service = (
            openai_service
        )


    # ======================================================
    # PROCESAR DOCUMENTO
    # ======================================================

    def translate_document(
        self,
        file_name: str,
        content_type: str,
        file_bytes: bytes
    ) -> dict:

        extension = (
            self._get_extension(
                file_name
            )
        )


        self._validate_file(
            extension=extension,
            content_type=content_type,
            file_bytes=file_bytes
        )


        # ==================================================
        # EXTRAER
        # ==================================================

        if (
            extension == "pdf"
        ):

            sections = (
                self._extract_pdf(
                    file_bytes
                )
            )


        elif (
            extension == "docx"
        ):

            sections = (
                self._extract_docx(
                    file_bytes
                )
            )


        elif (
            extension == "txt"
        ):

            sections = (
                self._extract_txt(
                    file_bytes
                )
            )


        else:

            raise DocumentServiceException(
                "Formato de documento no permitido.",
                415
            )


        # ==================================================
        # SIN CONTENIDO
        # ==================================================

        if (
            not sections
        ):

            raise DocumentServiceException(
                (
                    "No fue posible encontrar texto "
                    "legible dentro del documento."
                ),
                400
            )


        # ==================================================
        # CONTAR TEXTO
        # ==================================================

        character_count = sum(
            len(
                section["text"]
            )
            for section in sections
        )


        if (
            character_count == 0
        ):

            raise DocumentServiceException(
                (
                    "El documento no contiene "
                    "texto para traducir."
                ),
                400
            )


        if (
            character_count >
            self.MAX_EXTRACTED_CHARACTERS
        ):

            raise DocumentServiceException(
                (
                    "El documento contiene demasiado "
                    "texto para procesarlo en una sola "
                    "solicitud. Utiliza un documento "
                    "más corto."
                ),
                413
            )


        # ==================================================
        # OPENAI
        # ==================================================

        ai_result = (
            self._translate_sections(
                sections
            )
        )


        # ==================================================
        # COMBINAR RESULTADOS
        # ==================================================

        translated_items = {

            item["id"]:
                item

            for item in ai_result[
                "translations"
            ]
        }


        final_sections = []


        for section in sections:

            section_id = (
                section["id"]
            )


            translated = (
                translated_items.get(
                    section_id
                )
            )


            if (
                translated is None
            ):

                raise DocumentServiceException(
                    (
                        "La Inteligencia Artificial no "
                        "devolvió todas las secciones "
                        "del documento."
                    ),
                    502
                )


            final_sections.append({

                "id":
                    section_id,

                "title":
                    section["title"],

                "translated_title":
                    translated[
                        "translated_title"
                    ],

                "original_text":
                    section["text"],

                "translated_text":
                    translated[
                        "translated_text"
                    ]

            })


        return {

            "file_name":
                file_name,

            "file_type":
                extension.upper(),

            "file_size":
                len(
                    file_bytes
                ),

            "source_language":
                ai_result[
                    "source_language"
                ],

            "target_language":
                ai_result[
                    "target_language"
                ],

            "character_count":
                character_count,

            "section_count":
                len(
                    final_sections
                ),

            "sections":
                final_sections

        }


    # ======================================================
    # VALIDAR ARCHIVO
    # ======================================================

    def _validate_file(
        self,
        extension: str,
        content_type: str,
        file_bytes: bytes
    ):

        if (
            not file_bytes
        ):

            raise DocumentServiceException(
                "El archivo recibido está vacío.",
                400
            )


        if (
            extension
            not in self.ALLOWED_EXTENSIONS
        ):

            raise DocumentServiceException(
                (
                    "Formato no permitido. "
                    "Utiliza PDF, DOCX o TXT."
                ),
                415
            )


        max_bytes = (
            self.settings
            .max_document_size_mb
            * 1024
            * 1024
        )


        if (
            len(file_bytes) >
            max_bytes
        ):

            raise DocumentServiceException(
                (
                    "El documento supera el límite "
                    f"de {self.settings.max_document_size_mb} MB."
                ),
                413
            )


        clean_content_type = (
            content_type
            .split(";")[0]
            .strip()
            .lower()
        )


        allowed_types = (
            self.ALLOWED_MIME_TYPES[
                extension
            ]
        )


        if (
            clean_content_type
            and
            clean_content_type
            not in allowed_types
        ):

            raise DocumentServiceException(
                (
                    "El tipo de archivo recibido "
                    "no coincide con el formato permitido."
                ),
                415
            )


    # ======================================================
    # EXTENSIÓN
    # ======================================================

    @staticmethod
    def _get_extension(
        file_name: str
    ) -> str:

        if (
            not file_name or
            "." not in file_name
        ):

            return ""


        return (
            file_name
            .rsplit(
                ".",
                1
            )[1]
            .lower()
            .strip()
        )


    # ======================================================
    # PDF
    # ======================================================

    def _extract_pdf(
        self,
        file_bytes: bytes
    ) -> list[dict]:

        try:

            reader = PdfReader(
                io.BytesIO(
                    file_bytes
                )
            )


        except Exception:

            raise DocumentServiceException(
                (
                    "No fue posible leer el PDF. "
                    "El archivo puede estar dañado."
                ),
                400
            )


        if (
            reader.is_encrypted
        ):

            try:

                result = reader.decrypt(
                    ""
                )


                if (
                    result == 0
                ):

                    raise DocumentServiceException(
                        (
                            "El PDF está protegido "
                            "con contraseña."
                        ),
                        400
                    )


            except DocumentServiceException:

                raise


            except Exception:

                raise DocumentServiceException(
                    (
                        "No es posible procesar "
                        "un PDF protegido."
                    ),
                    400
                )


        sections = []


        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            try:

                text = (
                    page.extract_text()
                    or ""
                )


            except Exception:

                text = ""


            text = (
                self._clean_text(
                    text
                )
            )


            if (
                not text
            ):

                continue


            pieces = (
                self._split_large_text(
                    text
                )
            )


            for part_number, piece in enumerate(
                pieces,
                start=1
            ):

                title = (
                    f"Página {page_number}"
                )


                if (
                    len(pieces) > 1
                ):

                    title += (
                        f" · Parte {part_number}"
                    )


                sections.append({

                    "id":
                        len(sections) + 1,

                    "title":
                        title,

                    "text":
                        piece

                })


        return sections


    # ======================================================
    # DOCX
    # ======================================================

    def _extract_docx(
        self,
        file_bytes: bytes
    ) -> list[dict]:

        try:

            document = Document(
                io.BytesIO(
                    file_bytes
                )
            )


        except Exception:

            raise DocumentServiceException(
                (
                    "No fue posible leer el archivo DOCX. "
                    "El documento puede estar dañado."
                ),
                400
            )


        raw_sections = []


        current_title = (
            "Contenido"
        )


        current_paragraphs = []


        # ==================================================
        # PÁRRAFOS
        # ==================================================

        for paragraph in document.paragraphs:

            text = (
                self._clean_text(
                    paragraph.text
                )
            )


            if (
                not text
            ):

                continue


            style_name = (
                paragraph.style.name
                if paragraph.style
                else ""
            )


            is_heading = (
                style_name
                .lower()
                .startswith(
                    "heading"
                )
            )


            if (
                is_heading
            ):

                if (
                    current_paragraphs
                ):

                    raw_sections.append({

                        "title":
                            current_title,

                        "text":
                            "\n\n".join(
                                current_paragraphs
                            )

                    })


                    current_paragraphs = []


                current_title = (
                    text
                )

            else:

                current_paragraphs.append(
                    text
                )


        if (
            current_paragraphs
        ):

            raw_sections.append({

                "title":
                    current_title,

                "text":
                    "\n\n".join(
                        current_paragraphs
                    )

            })


        # ==================================================
        # TABLAS
        # ==================================================

        for table_number, table in enumerate(
            document.tables,
            start=1
        ):

            rows = []


            for row in table.rows:

                cells = [

                    self._clean_text(
                        cell.text
                    )

                    for cell in row.cells

                ]


                if (
                    any(cells)
                ):

                    rows.append(
                        " | ".join(
                            cells
                        )
                    )


            if (
                rows
            ):

                raw_sections.append({

                    "title":
                        f"Tabla {table_number}",

                    "text":
                        "\n".join(
                            rows
                        )

                })


        return (
            self._normalize_sections(
                raw_sections
            )
        )


    # ======================================================
    # TXT
    # ======================================================

    def _extract_txt(
        self,
        file_bytes: bytes
    ) -> list[dict]:

        text = None


        encodings = [
            "utf-8-sig",
            "utf-8",
            "latin-1"
        ]


        for encoding in encodings:

            try:

                text = file_bytes.decode(
                    encoding
                )

                break


            except UnicodeDecodeError:

                continue


        if (
            text is None
        ):

            raise DocumentServiceException(
                (
                    "No fue posible determinar la "
                    "codificación del archivo TXT."
                ),
                400
            )


        text = (
            self._clean_text(
                text
            )
        )


        if (
            not text
        ):

            return []


        blocks = re.split(
            r"\n\s*\n",
            text
        )


        raw_sections = []


        for number, block in enumerate(
            blocks,
            start=1
        ):

            block = (
                self._clean_text(
                    block
                )
            )


            if (
                not block
            ):

                continue


            raw_sections.append({

                "title":
                    f"Sección {number}",

                "text":
                    block

            })


        return (
            self._normalize_sections(
                raw_sections
            )
        )


    # ======================================================
    # NORMALIZAR SECCIONES
    # ======================================================

    def _normalize_sections(
        self,
        raw_sections: list[dict]
    ) -> list[dict]:

        sections = []


        for raw_section in raw_sections:

            title = (
                raw_section.get(
                    "title",
                    "Sección"
                )
            )


            text = (
                self._clean_text(
                    raw_section.get(
                        "text",
                        ""
                    )
                )
            )


            if (
                not text
            ):

                continue


            pieces = (
                self._split_large_text(
                    text
                )
            )


            for part_number, piece in enumerate(
                pieces,
                start=1
            ):

                final_title = (
                    title
                )


                if (
                    len(pieces) > 1
                ):

                    final_title += (
                        f" · Parte {part_number}"
                    )


                sections.append({

                    "id":
                        len(sections) + 1,

                    "title":
                        final_title,

                    "text":
                        piece

                })


        return sections


    # ======================================================
    # DIVIDIR TEXTO GRANDE
    # ======================================================

    def _split_large_text(
        self,
        text: str
    ) -> list[str]:

        if (
            len(text) <=
            self.MAX_SECTION_CHARACTERS
        ):

            return [
                text
            ]


        paragraphs = re.split(
            r"\n+",
            text
        )


        chunks = []

        current = []


        current_length = 0


        for paragraph in paragraphs:

            paragraph = (
                paragraph.strip()
            )


            if (
                not paragraph
            ):

                continue


            paragraph_length = (
                len(paragraph)
            )


            if (
                paragraph_length >
                self.MAX_SECTION_CHARACTERS
            ):

                if (
                    current
                ):

                    chunks.append(
                        "\n".join(
                            current
                        )
                    )

                    current = []

                    current_length = 0


                for start in range(
                    0,
                    paragraph_length,
                    self.MAX_SECTION_CHARACTERS
                ):

                    chunks.append(
                        paragraph[
                            start:
                            start +
                            self.MAX_SECTION_CHARACTERS
                        ]
                    )


                continue


            if (
                current_length +
                paragraph_length +
                1 >
                self.MAX_SECTION_CHARACTERS
            ):

                chunks.append(
                    "\n".join(
                        current
                    )
                )


                current = [
                    paragraph
                ]


                current_length = (
                    paragraph_length
                )


            else:

                current.append(
                    paragraph
                )


                current_length += (
                    paragraph_length + 1
                )


        if (
            current
        ):

            chunks.append(
                "\n".join(
                    current
                )
            )


        return chunks


    # ======================================================
    # LIMPIAR TEXTO
    # ======================================================

    @staticmethod
    def _clean_text(
        text: str
    ) -> str:

        text = (
            text.replace(
                "\x00",
                ""
            )
        )


        text = re.sub(
            r"[ \t]+",
            " ",
            text
        )


        text = re.sub(
            r"\n{3,}",
            "\n\n",
            text
        )


        return (
            text.strip()
        )


    # ======================================================
    # JSON SCHEMA PARA OPENAI
    # ======================================================

    @staticmethod
    def _translation_schema() -> dict:

        return {

            "type":
                "object",

            "properties": {

                "source_language": {

                    "type":
                        "string",

                    "enum": [
                        "es",
                        "en"
                    ]

                },

                "target_language": {

                    "type":
                        "string",

                    "enum": [
                        "es",
                        "en"
                    ]

                },

                "translations": {

                    "type":
                        "array",

                    "items": {

                        "type":
                            "object",

                        "properties": {

                            "id": {
                                "type":
                                    "integer"
                            },

                            "translated_title": {
                                "type":
                                    "string"
                            },

                            "translated_text": {
                                "type":
                                    "string"
                            }

                        },

                        "required": [
                            "id",
                            "translated_title",
                            "translated_text"
                        ],

                        "additionalProperties":
                            False

                    }

                }

            },

            "required": [
                "source_language",
                "target_language",
                "translations"
            ],

            "additionalProperties":
                False

        }


    # ======================================================
    # TRADUCIR CON OPENAI
    # ======================================================

    def _translate_sections(
        self,
        sections: list[dict]
    ) -> dict:

        instructions = """
You are a professional Spanish-English document translator.

You will receive structured document sections.

RULES:

1. Detect the predominant language of the document.

2. Only these source languages are valid:
   Spanish = es
   English = en

3. If the source language is Spanish:
   target_language must be "en".

4. If the source language is English:
   target_language must be "es".

5. Translate EVERY section completely.

6. Do not summarize.

7. Do not omit information.

8. Preserve names, numbers, dates, URLs, email addresses,
   currencies, units, acronyms and technical terms.

9. Preserve paragraph breaks whenever possible.

10. Translate section titles naturally.

11. Preserve table rows and separators when present.

12. Return exactly one translation entry for every input id.

13. Never invent content not present in the document.

14. Prefer natural professional translations over literal
    word-for-word translations.
""".strip()


        input_payload = {

            "sections": [

                {

                    "id":
                        section["id"],

                    "title":
                        section["title"],

                    "text":
                        section["text"]

                }

                for section in sections

            ]

        }


        result = (
            self.openai_service
            .create_structured_response(

                instructions=(
                    instructions
                ),

                input_text=(
                    json.dumps(
                        input_payload,
                        ensure_ascii=False
                    )
                ),

                schema_name=(
                    "translated_document"
                ),

                schema=(
                    self._translation_schema()
                ),

                max_output_tokens=12000

            )
        )


        # ==================================================
        # VALIDACIÓN IDIOMAS
        # ==================================================

        source_language = (
            result.get(
                "source_language"
            )
        )


        target_language = (
            result.get(
                "target_language"
            )
        )


        if (
            source_language == "es" and
            target_language != "en"
        ):

            raise DocumentServiceException(
                (
                    "La traducción devolvió una "
                    "dirección de idioma inválida."
                ),
                502
            )


        if (
            source_language == "en" and
            target_language != "es"
        ):

            raise DocumentServiceException(
                (
                    "La traducción devolvió una "
                    "dirección de idioma inválida."
                ),
                502
            )


        return result